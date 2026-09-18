import streamlit as st
from docx import Document
from docx.shared import Pt, Cm
import io

# Configuração básica da página
st.set_page_config(
    page_title="DocClean & Layout",
    page_icon="📄",
    layout="centered"
)

st.title("📄 DocClean & Layout")
st.write(
    "Bem-vindo! Este aplicativo vai ajudar a formatar automaticamente "
    "seus arquivos .docx (fontes, espaçamento, margens, etc.)."
)

st.divider()

# --- Opções de formatação escolhidas pelo usuário ---
st.subheader("⚙️ Opções de formatação")

col1, col2 = st.columns(2)
with col1:
    fonte_escolhida = st.selectbox("Fonte", ["Arial", "Times New Roman", "Calibri"])
with col2:
    tamanho_escolhido = st.selectbox("Tamanho da fonte", [10, 11, 12, 14], index=2)

espacamento_escolhido = st.selectbox("Espaçamento entre linhas", [1.0, 1.5, 2.0], index=1)
margem_escolhida = st.selectbox("Margem (cm)", [2.0, 2.5, 3.0], index=1)

st.divider()

# --- Upload do arquivo ---
arquivo = st.file_uploader("Envie seu arquivo .docx", type=["docx"])


def formatar_documento(documento, fonte, tamanho, espacamento, margem):
    """
    Recebe um objeto Document (python-docx) e aplica:
    - fonte e tamanho em todos os textos (runs)
    - espaçamento entre linhas em todos os parágrafos
    - margens em todas as seções do documento
    """
    # Aplica fonte e tamanho em cada trecho de texto (run) de cada parágrafo
    for paragrafo in documento.paragraphs:
        for run in paragrafo.runs:
            run.font.name = fonte
            run.font.size = Pt(tamanho)

    # Aplica espaçamento entre linhas em cada parágrafo
    for paragrafo in documento.paragraphs:
        paragrafo.paragraph_format.line_spacing = espacamento

    # Aplica margens em todas as seções do documento
    for secao in documento.sections:
        secao.top_margin = Cm(margem)
        secao.bottom_margin = Cm(margem)
        secao.left_margin = Cm(margem)
        secao.right_margin = Cm(margem)

    return documento


if arquivo is not None:
    st.success(f"Arquivo recebido: {arquivo.name}")

    documento = Document(io.BytesIO(arquivo.getvalue()))

    total_paragrafos = len(documento.paragraphs)

    preview_texto = ""
    contador = 0
    for paragrafo in documento.paragraphs:
        if paragrafo.text.strip() != "":
            preview_texto += paragrafo.text + "\n\n"
            contador += 1
        if contador >= 3:
            break

    st.divider()
    st.subheader("📊 Informações do documento original")
    st.write(f"**Total de parágrafos:** {total_paragrafos}")
    st.subheader("👀 Prévia do texto original")
    st.text(preview_texto if preview_texto else "O documento não contém texto visível.")

    st.divider()

    # Botão para acionar a formatação
    if st.button("✨ Formatar documento"):
        documento_formatado = formatar_documento(
            documento,
            fonte_escolhida,
            tamanho_escolhido,
            espacamento_escolhido,
            margem_escolhida
        )

        # Salva o documento formatado em memória (não no disco)
        buffer_saida = io.BytesIO()
        documento_formatado.save(buffer_saida)
        buffer_saida.seek(0)

        st.success("Documento formatado com sucesso!")

        st.download_button(
            label="⬇️ Baixar documento formatado",
            data=buffer_saida,
            file_name=f"formatado_{arquivo.name}",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

else:
    st.warning("Nenhum arquivo enviado ainda.")

