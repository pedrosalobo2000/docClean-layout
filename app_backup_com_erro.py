184
import streamlit as st
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io

# Configuração básica da página
st.set_page_config(
    page_title="DocClean & Layout",
    page_icon="📄",
    layout="centered"
)

st.title("📄 DocClean & Layout")
st.write(
    "Bem-vindo! Este aplicativo vai ajudar a formatar automaticamente "
    "seus arquivos .docx (fontes, espaçamento, margens, etc.)."
)

st.divider()

# --- Opções de formatação escolhidas pelo usuário ---
st.subheader("⚙️ Opções de formatação")

usar_padrao_abnt = st.checkbox("📐 Aplicar padrão ABNT completo (recomendado para trabalhos acadêmicos)")

if usar_padrao_abnt:
    st.info(
        "Padrão ABNT selecionado: Times New Roman 12, espaçamento 1.5, "
        "margens 3cm (superior/esquerda) e 2cm (inferior/direita)."
    )
    fonte_escolhida = "Times New Roman"
    tamanho_escolhido = 12
    espacamento_escolhido = 1.5
else:
    col1, col2 = st.columns(2)
    with col1:
        fonte_escolhida = st.selectbox("Fonte", ["Arial", "Times New Roman", "Calibri"])
    with col2:
        tamanho_escolhido = st.selectbox("Tamanho da fonte", [10, 11, 12, 14], index=2)

    espacamento_escolhido = st.selectbox("Espaçamento entre linhas", [1.0, 1.5, 2.0], index=1)
    margem_escolhida = st.selectbox("Margem (cm)", [2.0, 2.5, 3.0], index=1)

# --- Upload do arquivo ---
arquivo = st.file_uploader("Envie seu arquivo .docx", type=["docx"])


# ============================================================
# REGRAS ABNT (valores de referência para o verificador)
# ============================================================
ABNT_FONTES_VALIDAS = ["Arial", "Times New Roman"]
ABNT_TAMANHO_FONTE = 12
ABNT_ESPACAMENTO = 1.5
ABNT_MARGEM_SUPERIOR_CM = 3.0
ABNT_MARGEM_ESQUERDA_CM = 3.0
ABNT_MARGEM_INFERIOR_CM = 2.0
ABNT_MARGEM_DIREITA_CM = 2.0


def verificar_abnt(documento):
    """
    Analisa o documento e retorna uma lista de problemas encontrados,
    comparando com as regras ABNT definidas acima.
    Cada problema é um dicionário com: tipo, local e mensagem.
    """
    problemas = []

    # --- 1. Verificar margens (a regra vale para o documento todo) ---
    for secao in documento.sections:
        margem_top_cm = secao.top_margin.cm
        margem_left_cm = secao.left_margin.cm
        margem_bottom_cm = secao.bottom_margin.cm
        margem_right_cm = secao.right_margin.cm

        if round(margem_top_cm, 1) != ABNT_MARGEM_SUPERIOR_CM:
            problemas.append({
                "tipo": "Margem",
                "local": "Página",
                "mensagem": f"Margem superior é {margem_top_cm:.1f}cm (ABNT exige {ABNT_MARGEM_SUPERIOR_CM}cm)"
            })
        if round(margem_left_cm, 1) != ABNT_MARGEM_ESQUERDA_CM:
            problemas.append({
                "tipo": "Margem",
                "local": "Página",
                "mensagem": f"Margem esquerda é {margem_left_cm:.1f}cm (ABNT exige {ABNT_MARGEM_ESQUERDA_CM}cm)"
            })
        if round(margem_bottom_cm, 1) != ABNT_MARGEM_INFERIOR_CM:
            problemas.append({
                "tipo": "Margem",
                "local": "Página",
                "mensagem": f"Margem inferior é {margem_bottom_cm:.1f}cm (ABNT exige {ABNT_MARGEM_INFERIOR_CM}cm)"
            })
        if round(margem_right_cm, 1) != ABNT_MARGEM_DIREITA_CM:
            problemas.append({
                "tipo": "Margem",
                "local": "Página",
                "mensagem": f"Margem direita é {margem_right_cm:.1f}cm (ABNT exige {ABNT_MARGEM_DIREITA_CM}cm)"
            })

    # --- 2. Verificar cada parágrafo: fonte, tamanho, espaçamento e alinhamento ---
    for numero, paragrafo in enumerate(documento.paragraphs, start=1):
        # Pula parágrafos vazios (não têm o que verificar)
        if paragrafo.text.strip() == "":
            continue

        # --- Espaçamento entre linhas ---
        espacamento = paragrafo.paragraph_format.line_spacing
        if espacamento is not None and round(espacamento, 1) != ABNT_ESPACAMENTO:
            problemas.append({
                "tipo": "Espaçamento",
                "local": f"Parágrafo {numero}",
                "mensagem": f"Espaçamento é {espacamento} (ABNT exige {ABNT_ESPACAMENTO})"
            })

        # --- Alinhamento (ABNT exige texto justificado no corpo do trabalho) ---
                # --- Alinhamento (ABNT exige texto justificado no corpo do trabalho) ---
        try:
            alinhamento = paragrafo.alignment
        except ValueError:
            # Alguns arquivos .docx (geralmente convertidos de outros formatos)
            # têm um valor de alinhamento que o python-docx não reconhece.
            # Nesse caso, tratamos como "não foi possível verificar" e seguimos.
            alinhamento = None

        if alinhamento is not None and alinhamento != WD_ALIGN_PARAGRAPH.JUSTIFY:
            problemas.append({
                "tipo": "Alinhamento",
                "local": f"Parágrafo {numero}",
                "mensagem": "Parágrafo não está justificado"
            })

        # --- Fonte e tamanho (verifica cada trecho de texto do parágrafo) ---
        for run in paragrafo.runs:
            if run.text.strip() == "":
                continue

            nome_fonte = run.font.name
            if nome_fonte is not None and nome_fonte not in ABNT_FONTES_VALIDAS:
                problemas.append({
                    "tipo": "Fonte",
                    "local": f"Parágrafo {numero}",
                    "mensagem": f"Fonte '{nome_fonte}' não é padrão ABNT (use Arial ou Times New Roman)"
                })

            tamanho_fonte = run.font.size
            if tamanho_fonte is not None and tamanho_fonte.pt != ABNT_TAMANHO_FONTE:
                problemas.append({
                    "tipo": "Tamanho da fonte",
                    "local": f"Parágrafo {numero}",
                    "mensagem": f"Tamanho é {tamanho_fonte.pt}pt (ABNT exige {ABNT_TAMANHO_FONTE}pt)"
                })

    return problemas


def formatar_documento(documento, fonte, tamanho, espacamento,
                        margem_superior, margem_inferior, margem_esquerda, margem_direita):
    """
    Recebe um objeto Document (python-docx) e aplica:
    - fonte e tamanho em todos os textos (runs)
    - espaçamento entre linhas em todos os parágrafos
    - alinhamento justificado em todos os parágrafos
    - margens em todas as seções do documento (cada lado pode ter um valor diferente)
    """
    for paragrafo in documento.paragraphs:
        for run in paragrafo.runs:
            run.font.name = fonte
            run.font.size = Pt(tamanho)

    for paragrafo in documento.paragraphs:
        paragrafo.paragraph_format.line_spacing = espacamento
        if paragrafo.text.strip() != "":
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for secao in documento.sections:
        secao.top_margin = Cm(margem_superior)
        secao.bottom_margin = Cm(margem_inferior)
        secao.left_margin = Cm(margem_esquerda)
        secao.right_margin = Cm(margem_direita)

    return documento
    for paragrafo in do

if arquivo is not None:
    st.success(f"Arquivo recebido: {arquivo.name}")

    documento = Document(io.BytesIO(arquivo.getvalue()))

    total_paragrafos = len(documento.paragraphs)

    preview_texto = ""
    contador = 0
    for paragrafo in documento.paragraphs:
        if paragrafo.text.strip() != "":
            preview_texto += paragrafo.text + "\n\n"
            contador += 1
        if contador >= 3:
            break

    st.divider()
    st.subheader("📊 Informações do documento original")
    st.write(f"**Total de parágrafos:** {total_paragrafos}")
    st.subheader("👀 Prévia do texto original")
    st.text(preview_texto if preview_texto else "O documento não contém texto visível.")

    st.divider()

    # --- Botão do verificador ABNT ---
    st.subheader("🔍 Verificador de conformidade ABNT")
    if st.button("Verificar conformidade ABNT"):
        problemas = verificar_abnt(documento)

        if not problemas:
            st.success("✅ Nenhum problema encontrado! O documento segue as regras ABNT verificadas.")
        else:
            st.warning(f"⚠️ Foram encontrados {len(problemas)} problema(s):")
            for problema in problemas:
                st.write(f"**[{problema['tipo']}]** {problema['local']}: {problema['mensagem']}")

    st.divider()

    # --- Botão de formatação automática ---
    st.subheader("✨ Formatação automática")
        if st.button("Formatar documento"):
        if usar_padrao_abnt:
            margem_sup, margem_inf, margem_esq, margem_dir = 3.0, 2.0, 3.0, 2.0
        else:
            margem_sup, margem_inf, margem_esq, margem_dir = (
                margem_escolhida, margem_escolhida, margem_escolhida, margem_escolhida
            )

        documento_formatado = formatar_documento(
            documento,
            fonte_escolhida,
            tamanho_escolhido,
            espacamento_escolhido,
            margem_sup, margem_inf, margem_esq, margem_dir
        )

else:
    st.warning("Nenhum arquivo enviado ainda.")

